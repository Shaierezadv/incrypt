import io
from pathlib import Path
from typing import Tuple


def read_text_auto(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8")
    if ext == ".docx":
        try:
            from docx import Document  # type: ignore
        except Exception as e:
            raise RuntimeError("python-docx is required to read .docx files. Install with: pip install python-docx")
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        return "\n".join(lines)
    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception as e:
            raise RuntimeError("pypdf is required to read .pdf files. Install with: pip install pypdf")
        reader = PdfReader(str(path))
        texts = []
        for page in reader.pages:
            txt = page.extract_text() or ""
            texts.append(txt)
        return "\n".join(texts)
    raise ValueError(f"Unsupported file extension: {ext}")


def write_text_auto(path: Path, text: str) -> None:
    ext = path.suffix.lower()
    if ext == ".txt":
        path.write_text(text, encoding="utf-8")
        return
    if ext == ".docx":
        try:
            from docx import Document  # type: ignore
        except Exception as e:
            raise RuntimeError("python-docx is required to write .docx files. Install with: pip install python-docx")
        doc = Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(str(path))
        return
    if ext == ".pdf":
        try:
            from reportlab.pdfgen import canvas  # type: ignore
            from reportlab.lib.pagesizes import A4  # type: ignore
            from reportlab.lib.units import mm  # type: ignore
        except Exception:
            raise RuntimeError("reportlab is required to write .pdf files. Install with: pip install reportlab")
        c = canvas.Canvas(str(path), pagesize=A4)
        width, height = A4
        left_margin = 20 * mm
        right_margin = 20 * mm
        top_margin = 20 * mm
        bottom_margin = 20 * mm
        usable_width = width - left_margin - right_margin
        line_height = 14
        x = left_margin
        y = height - top_margin
        c.setFont("Helvetica", 11)

        def draw_line(s: str):
            nonlocal x, y
            if y < bottom_margin:
                c.showPage()
                c.setFont("Helvetica", 11)
                y = height - top_margin
            c.drawString(x, y, s)
            y -= line_height

        for line in text.splitlines():
            # naive wrap
            remaining = line
            while remaining:
                # estimate chars per line by width/font size (roughly 0.5*font_size per char in Helvetica)
                max_chars = int(usable_width / 6.0)
                seg = remaining[:max_chars]
                draw_line(seg)
                remaining = remaining[max_chars:]
        c.save()
        return
    raise ValueError(f"Unsupported file extension: {ext}")